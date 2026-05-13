from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# ── Palette ───────────────────────────────────────────────────────────────────
NAVY    = RGBColor(0x0D, 0x1B, 0x2A)
TEAL    = RGBColor(0x00, 0x8B, 0x8B)
SLATE   = RGBColor(0x44, 0x55, 0x66)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT   = RGBColor(0xF4, 0xF6, 0xF9)
ACCENT  = RGBColor(0x00, 0xC2, 0xB2)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)


def rgb_fill(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def add_textbox(slide, text, left, top, width, height,
                font_size=11, bold=False, color=WHITE,
                align=PP_ALIGN.LEFT, wrap=True):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return txBox


def add_bullet_box(slide, items, left, top, width, height,
                   font_size=10.5, heading=None, heading_color=ACCENT,
                   text_color=WHITE, line_spacing=None):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    first = True
    if heading:
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = heading
        run.font.size = Pt(font_size + 1)
        run.font.bold = True
        run.font.color.rgb = heading_color
        first = False

    for item in items:
        p = tf.add_paragraph() if not first else tf.paragraphs[0]
        first = False
        p.space_before = Pt(3)
        if line_spacing:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = item
        run.font.size = Pt(font_size)
        run.font.color.rgb = text_color
    return txBox


def build():
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H

    blank = prs.slide_layouts[6]  # completely blank

    # ══════════════════════════════════════════════════════════════════════════
    # SLIDE 1 — Title / Overview
    # ══════════════════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank)

    # Full background
    bg = s1.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    rgb_fill(bg, NAVY)
    bg.line.fill.background()

    # Left accent bar
    bar = s1.shapes.add_shape(1, 0, 0, Inches(0.18), SLIDE_H)
    rgb_fill(bar, TEAL)
    bar.line.fill.background()

    # App name
    add_textbox(s1, "Trinket Intelligence",
                Inches(0.45), Inches(1.6), Inches(8), Inches(1.1),
                font_size=40, bold=True, color=WHITE)

    # Subtitle
    add_textbox(s1, "Methodology & Data Sources",
                Inches(0.45), Inches(2.65), Inches(8), Inches(0.6),
                font_size=18, bold=False, color=ACCENT)

    # Divider line
    line = s1.shapes.add_shape(1, Inches(0.45), Inches(3.35), Inches(5.5), Inches(0.04))
    rgb_fill(line, TEAL)
    line.line.fill.background()

    # Description paragraph
    desc = (
        "Trinket Intelligence is a sandbox application that allows users to look up "
        "any publicly listed company by ticker symbol and retrieve real-time market data, "
        "industry peer comparisons, and the latest news headlines — all sourced on demand "
        "from Yahoo Finance's publicly available APIs."
    )
    add_textbox(s1, desc,
                Inches(0.45), Inches(3.55), Inches(7.6), Inches(1.6),
                font_size=12, color=RGBColor(0xCC, 0xD6, 0xE0))

    # Footer note
    add_textbox(s1, "A sandbox / experiment project  ·  Data provided by Yahoo Finance",
                Inches(0.45), Inches(6.7), Inches(9), Inches(0.5),
                font_size=9, color=RGBColor(0x66, 0x77, 0x88))

    # Right-side pill badges
    badges = [
        ("Market Data",    Inches(9.8), Inches(1.9)),
        ("Industry Peers", Inches(9.8), Inches(2.7)),
        ("News Headlines", Inches(9.8), Inches(3.5)),
    ]
    for label, bx, by in badges:
        pill = s1.shapes.add_shape(1, bx, by, Inches(2.8), Inches(0.52))
        rgb_fill(pill, TEAL)
        pill.line.fill.background()
        add_textbox(s1, label, bx, by + Inches(0.06), Inches(2.8), Inches(0.42),
                    font_size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # ══════════════════════════════════════════════════════════════════════════
    # SLIDE 2 — Methodology detail (3-column)
    # ══════════════════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank)

    bg2 = s2.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    rgb_fill(bg2, NAVY)
    bg2.line.fill.background()

    bar2 = s2.shapes.add_shape(1, 0, 0, Inches(0.18), SLIDE_H)
    rgb_fill(bar2, TEAL)
    bar2.line.fill.background()

    # Slide title
    add_textbox(s2, "How It Works",
                Inches(0.45), Inches(0.28), Inches(10), Inches(0.65),
                font_size=26, bold=True, color=WHITE)

    line2 = s2.shapes.add_shape(1, Inches(0.45), Inches(0.95), Inches(12.4), Inches(0.04))
    rgb_fill(line2, TEAL)
    line2.line.fill.background()

    # ── Column layout ─────────────────────────────────────────────────────────
    col_top    = Inches(1.15)
    col_h      = Inches(5.6)
    col_w      = Inches(3.8)
    col_pad    = Inches(0.26)
    col_starts = [Inches(0.45), Inches(4.55), Inches(8.65)]

    sections = [
        {
            "title": "1  Market Data",
            "source": "Source: Yahoo Finance  /  yfinance API",
            "bullets": [
                "• Fetched in real time when a ticker is searched",
                "• Current price & market capitalisation",
                "• Trailing P/E ratio",
                "• 52-week high and low",
                "• Daily volume & 3-month average volume",
                "• Sector and industry classification",
                "",
                "Data is cached for 5 minutes to limit redundant\n"
                "API calls on repeated lookups.",
            ],
        },
        {
            "title": "2  Industry Peers",
            "source": "Source: Yahoo Finance Equity Screener",
            "bullets": [
                "• The searched company's industry label is\n"
                "  retrieved from Yahoo Finance (e.g. Microsoft\n"
                "  → Software—Infrastructure)",
                "• Yahoo Finance's screener finds all companies\n"
                "  sharing that industry classification",
                "• Filtered to US-listed stocks only\n"
                "  (NYSE & NASDAQ)",
                "• Sorted by market capitalisation, largest first",
                "• Preferred shares & depositary units excluded",
                "• Up to 10 peers displayed",
                "",
                "Limitation: peers reflect Yahoo Finance's own\n"
                "taxonomy. Conglomerates may show unexpected\n"
                "groupings based on their primary classification.",
            ],
        },
        {
            "title": "3  News Headlines",
            "source": "Source: Yahoo Finance Search API",
            "bullets": [
                "• Uses the yfinance Search component, which\n"
                "  queries Yahoo Finance's news index",
                "• Returns the 3 most recent articles\n"
                "  associated with the ticker symbol",
                "• Each headline links directly to the\n"
                "  original source article",
                "• Publisher name and publish date shown\n"
                "  where available",
                "",
                "Headlines reflect Yahoo Finance's news feed\n"
                "and may include press releases, analyst notes,\n"
                "and third-party editorial content.",
            ],
        },
    ]

    for i, sec in enumerate(sections):
        cx = col_starts[i]

        # Card background
        card = s2.shapes.add_shape(1, cx, col_top, col_w, col_h)
        rgb_fill(card, RGBColor(0x14, 0x28, 0x3C))
        card.line.color.rgb = TEAL
        card.line.width = Pt(0.75)

        # Top colour band
        band = s2.shapes.add_shape(1, cx, col_top, col_w, Inches(0.08))
        rgb_fill(band, TEAL)
        band.line.fill.background()

        # Section title
        add_textbox(s2, sec["title"],
                    cx + col_pad, col_top + Inches(0.18),
                    col_w - col_pad, Inches(0.55),
                    font_size=14, bold=True, color=WHITE)

        # Source line
        add_textbox(s2, sec["source"],
                    cx + col_pad, col_top + Inches(0.72),
                    col_w - col_pad, Inches(0.38),
                    font_size=8.5, bold=False, color=ACCENT)

        # Divider
        div = s2.shapes.add_shape(1,
                                  cx + col_pad,
                                  col_top + Inches(1.1),
                                  col_w - col_pad * 2,
                                  Inches(0.03))
        rgb_fill(div, RGBColor(0x1E, 0x3A, 0x52))
        div.line.fill.background()

        # Bullets
        add_bullet_box(s2, sec["bullets"],
                       cx + col_pad,
                       col_top + Inches(1.18),
                       col_w - col_pad,
                       col_h - Inches(1.3),
                       font_size=9.5,
                       text_color=RGBColor(0xCC, 0xD6, 0xE0))

    # Footer
    add_textbox(s2,
                "Trinket Intelligence  ·  A sandbox / experiment project  ·  "
                "All data sourced from Yahoo Finance's publicly available APIs  ·  "
                "Not intended for investment decisions",
                Inches(0.45), Inches(7.05), Inches(12.4), Inches(0.38),
                font_size=7.5, color=RGBColor(0x55, 0x66, 0x77))

    out = "Trinket_Intelligence_Methodology.pptx"
    prs.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    build()
